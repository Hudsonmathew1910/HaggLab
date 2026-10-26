"""
Generate the sample (fictional) datasets for each project's Try it page.

Run:  python datasets/generate.py
Output is the same on every run (fixed random seed). All companies, products,
customers and numbers are made up. No personal data is included.
"""
import csv
import math
import random
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent
rng = random.Random(2026)


def write_csv(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    print(f'{path.relative_to(ROOT)}: {len(rows)} rows')


def daterange(start, end):
    d = start
    while d <= end:
        yield d
        d += timedelta(days=1)


def seasonal(d):
    """Shopping pattern: November/December peak, quiet February, busier weekends."""
    month_factor = {1: 0.85, 2: 0.78, 3: 0.92, 4: 0.95, 5: 1.0, 6: 0.97, 7: 0.93,
                    8: 0.98, 9: 1.02, 10: 1.08, 11: 1.32, 12: 1.55}[d.month]
    weekday_factor = [0.9, 0.88, 0.92, 0.97, 1.12, 1.35, 1.2][d.weekday()]
    return month_factor * weekday_factor


# ---------------------------------------------------------------------------
# 1. InsightDesk: retail shop sales (18 months of transaction lines)
# ---------------------------------------------------------------------------
RETAIL_PRODUCTS = [
    # sku, product, category, price
    ('AUD-101', 'Studio Over-Ear Headphones', 'Audio', 149.00),
    ('AUD-102', 'Noise-Cancelling Earbuds', 'Audio', 129.00),
    ('AUD-103', 'Portable Waterproof Speaker', 'Audio', 79.00),
    ('AUD-104', 'USB Desk Microphone', 'Audio', 99.00),
    ('AUD-105', 'Kids Volume-Safe Headphones', 'Audio', 35.00),
    ('CMP-201', 'Ergonomic Wireless Keyboard', 'Computer Accessories', 89.00),
    ('CMP-202', 'Precision Wireless Mouse', 'Computer Accessories', 49.00),
    ('CMP-203', 'USB-C 7-in-1 Hub', 'Computer Accessories', 59.00),
    ('CMP-204', '1080p Webcam', 'Computer Accessories', 69.00),
    ('CMP-205', 'Laptop Stand Aluminium', 'Computer Accessories', 45.00),
    ('CMP-206', '1TB Portable SSD', 'Computer Accessories', 119.00),
    ('OFF-301', 'Mesh Office Chair', 'Office Furniture', 249.00),
    ('OFF-302', 'Standing Desk Frame', 'Office Furniture', 399.00),
    ('OFF-303', 'LED Desk Lamp', 'Office Furniture', 39.00),
    ('OFF-304', 'Monitor Arm Dual', 'Office Furniture', 129.00),
    ('OFF-305', 'Under-Desk Cable Tray', 'Office Furniture', 25.00),
    ('APP-401', 'Merino Crew Sweater', 'Apparel', 79.00),
    ('APP-402', 'Waterproof Rain Jacket', 'Apparel', 119.00),
    ('APP-403', 'Everyday Cotton T-Shirt', 'Apparel', 22.00),
    ('APP-404', 'Running Shoes Lightweight', 'Apparel', 95.00),
    ('APP-405', 'Thermal Gloves', 'Apparel', 19.00),
    ('HOM-501', 'Insulated Steel Bottle', 'Home & Kitchen', 29.00),
    ('HOM-502', 'Pour-Over Coffee Set', 'Home & Kitchen', 49.00),
    ('HOM-503', 'Cast Iron Skillet 26cm', 'Home & Kitchen', 55.00),
    ('HOM-504', 'Bamboo Cutting Board Set', 'Home & Kitchen', 32.00),
    ('FIT-601', 'Yoga Mat 6mm', 'Fitness', 35.00),
    ('FIT-602', 'Adjustable Dumbbell Pair', 'Fitness', 189.00),
    ('FIT-603', 'Resistance Band Set', 'Fitness', 24.00),
    ('FIT-604', 'Fitness Tracker Band', 'Fitness', 79.00),
]
# What each product costs the shop to buy, as a share of its list price (no random draws, so
# adding costs does not change any other number in the file). A few products have thin margins,
# so promotion discounts or the spring clearance push them below cost: real losses to find.
CATEGORY_COST = {'Audio': 0.62, 'Computer Accessories': 0.58, 'Office Furniture': 0.66,
                 'Apparel': 0.45, 'Home & Kitchen': 0.50, 'Fitness': 0.55}
COST_OVERRIDE = {'CMP-206': 0.88, 'AUD-105': 0.91, 'FIT-602': 0.87, 'APP-405': 0.70}
CLEARANCE = {'APP-405': (3, 4, 5)}  # Thermal Gloves: 40% off in Mar-May to clear winter stock


def unit_cost(sku, category, price):
    return round(price * COST_OVERRIDE.get(sku, CATEGORY_COST[category]), 2)


STORES = [('Downtown', 'North'), ('Riverside', 'South'), ('Harbour', 'East'), ('Hillview', 'West'), ('Online Store', 'Online')]
STORE_WEIGHT = [1.15, 0.95, 0.85, 0.75, 1.3]
# Product popularity, with a winter bump for apparel and a January bump for fitness
POPULARITY = {sku: rng.uniform(0.5, 1.6) for sku, *_ in RETAIL_PRODUCTS}


def product_weight(sku, category, d):
    w = POPULARITY[sku]
    if category == 'Apparel' and d.month in (10, 11, 12, 1):
        w *= 1.6
    if category == 'Fitness' and d.month == 1:
        w *= 1.9
    if category == 'Office Furniture' and d.month in (8, 9):
        w *= 1.4  # back-to-work season
    return w


def make_retail_sales():
    rows = []
    order_no = 100000
    start, end = date(2025, 4, 1), date(2026, 9, 30)
    promo_weeks = {date(2025, 11, 24), date(2025, 12, 15), date(2026, 3, 9), date(2026, 7, 13)}
    for d in daterange(start, end):
        week_start = d - timedelta(days=d.weekday())
        promo = week_start in promo_weeks
        growth = 1 + 0.12 * ((d - start).days / 365)  # the business grows about 12% a year
        for (store, region), sw in zip(STORES, STORE_WEIGHT):
            orders_today = max(1, int(rng.gauss(4.2, 1.2) * seasonal(d) * sw * growth * (1.3 if promo else 1)))
            for _ in range(orders_today):
                order_no += 1
                lines = rng.choices([1, 2, 3], weights=[0.62, 0.28, 0.10])[0]
                weights = [product_weight(s, c, d) for s, _, c, _ in RETAIL_PRODUCTS]
                chosen = {p[0]: p for p in rng.choices(RETAIL_PRODUCTS, weights=weights, k=lines)}
                payment = rng.choices(['Card', 'UPI', 'Cash', 'Wallet'],
                                      weights=[0.5, 0.3, 0.12 if store != 'Online Store' else 0, 0.08])[0]
                for sku, product, category, price in chosen.values():
                    qty = rng.choices([1, 2, 3], weights=[0.8, 0.16, 0.04])[0]
                    discount = 15 if promo else rng.choices([0, 5, 10], weights=[0.85, 0.1, 0.05])[0]
                    if d.month in CLEARANCE.get(sku, ()):
                        discount = 40
                    revenue = round(qty * price * (1 - discount / 100), 2)
                    cost = round(qty * unit_cost(sku, category, price), 2)
                    rows.append([d.isoformat(), f'ORD-{order_no}', store, region,
                                 'Online' if store == 'Online Store' else 'In-store',
                                 category, product, sku, qty, f'{price:.2f}', discount, f'{revenue:.2f}',
                                 f'{unit_cost(sku, category, price):.2f}', f'{cost:.2f}', f'{revenue - cost:.2f}', payment])
    write_csv(ROOT / 'insightdesk' / 'retail_sales.csv',
              ['date', 'order_id', 'store', 'region', 'channel', 'category', 'product', 'sku',
               'quantity', 'unit_price', 'discount_pct', 'revenue', 'unit_cost', 'cost', 'profit', 'payment_method'], rows)


# ---------------------------------------------------------------------------
# 2. ShopAssist: product catalog with real descriptions
# ---------------------------------------------------------------------------
CATALOG_SPEC = {
    'Headphones': ('Audio', ['Aurio', 'SonicPeak', 'Calmwave'], (39, 299), [
        ('Over-Ear Wireless Headphones', 'Closed-back over-ear headphones with {hours}-hour battery life and soft memory-foam cushions.', ['Bluetooth 5.3', 'Foldable design', 'USB-C fast charging']),
        ('Noise-Cancelling Headphones', 'Active noise cancelling blocks office and travel noise so you can focus. {hours}-hour battery.', ['Active noise cancelling', 'Transparency mode', 'Multipoint pairing']),
        ('Wired Studio Headphones', 'Flat, accurate sound for editing, mixing and long listening sessions.', ['Detachable 3m cable', '50mm drivers', 'Replaceable ear pads']),
        ('Kids Headphones', 'Volume limited to 85 dB to protect young ears, with a bendable, tough headband.', ['Volume limit 85 dB', 'Sharing port', 'Ages 3 and up']),
    ]),
    'Earbuds': ('Audio', ['Aurio', 'SonicPeak', 'Pulse'], (29, 199), [
        ('True Wireless Earbuds', 'Small, light earbuds with a pocket charging case that adds {hours} hours of play.', ['Touch controls', 'IPX4 sweat resistant', 'Wireless charging case']),
        ('Sport Earbuds', 'Secure ear hooks that stay in during runs and workouts.', ['IPX7 waterproof', 'Ear hooks', 'Quick pairing']),
        ('Noise-Cancelling Earbuds', 'Active noise cancelling in a compact earbud, great for commutes and flights.', ['Active noise cancelling', '3 ear-tip sizes', 'Low-latency mode']),
    ]),
    'Speakers': ('Audio', ['Boomly', 'SonicPeak'], (29, 249), [
        ('Portable Bluetooth Speaker', 'Room-filling sound in a speaker that fits in a backpack. {hours}-hour battery.', ['IP67 waterproof', 'Pair two speakers', 'Built-in microphone']),
        ('Smart Home Speaker', 'Voice-controlled speaker for music, timers and smart-home control.', ['Wi-Fi and Bluetooth', 'Voice assistant', 'Multi-room audio']),
        ('Desktop Speakers', 'Compact speakers for your computer with clear vocals and deep bass.', ['USB powered', 'Headphone jack', 'Volume knob']),
    ]),
    'Keyboards': ('Computer Accessories', ['KeyCraft', 'Typewise'], (25, 189), [
        ('Ergonomic Split Keyboard', 'Split, tented layout that keeps wrists straight during long typing days.', ['Wireless', 'Palm rest included', 'Mac and Windows keys']),
        ('Mechanical Keyboard', 'Tactile mechanical switches with a satisfying, accurate feel.', ['Hot-swappable switches', 'Backlit keys', 'Detachable USB-C cable']),
        ('Slim Wireless Keyboard', 'Quiet, low-profile keys and a rechargeable battery that lasts weeks.', ['Connects to 3 devices', 'Quiet keys', 'Rechargeable']),
    ]),
    'Mice': ('Computer Accessories', ['KeyCraft', 'Glide'], (15, 99), [
        ('Ergonomic Vertical Mouse', 'Vertical grip that reduces wrist strain.', ['Wireless', 'Adjustable DPI', 'Thumb buttons']),
        ('Silent Wireless Mouse', 'Silent clicks for shared offices and libraries.', ['Silent clicks', '18-month battery', 'Plug-and-play receiver']),
        ('Gaming Mouse', 'Fast, precise sensor with programmable buttons.', ['26,000 DPI sensor', '6 programmable buttons', 'RGB lighting']),
    ]),
    'Webcams': ('Computer Accessories', ['ClearCam'], (39, 159), [
        ('1080p Webcam', 'Sharp video calls with automatic light correction.', ['1080p at 30 fps', 'Privacy shutter', 'Dual microphones']),
        ('4K Webcam', 'Studio-quality 4K video with autofocus for streaming and meetings.', ['4K at 30 fps', 'Autofocus', 'Wide 90-degree view']),
    ]),
    'Chairs': ('Office Furniture', ['Sitwell', 'Posture Co'], (129, 549), [
        ('Mesh Office Chair', 'Breathable mesh back with adjustable lumbar support for all-day comfort.', ['Adjustable lumbar support', '3D armrests', 'Weight limit 130 kg']),
        ('Executive Leather Chair', 'Padded leather chair with a high back and headrest.', ['Headrest', 'Tilt lock', 'Weight limit 150 kg']),
        ('Kneeling Chair', 'Opens the hip angle to encourage an upright posture.', ['Wooden frame', 'Washable cushions', 'Folds flat']),
    ]),
    'Desks': ('Office Furniture', ['Sitwell', 'Workbase'], (149, 799), [
        ('Electric Standing Desk', 'Switch between sitting and standing at the touch of a button.', ['Height 71-121 cm', '3 memory presets', 'Anti-collision']),
        ('Compact Writing Desk', 'A small desk that fits bedrooms and small home offices.', ['Width {width} cm', 'Cable hole', 'Easy assembly']),
        ('Corner Desk', 'L-shaped desk with room for two monitors.', ['Width 140 x 120 cm', 'Steel frame', 'Reversible layout']),
    ]),
    'Jackets': ('Apparel', ['Northtrail', 'Urbano'], (59, 249), [
        ('Waterproof Rain Jacket', 'Fully taped seams and a breathable shell keep you dry in heavy rain.', ['Waterproof 10,000 mm', 'Packable hood', 'Sizes XS-XXL']),
        ('Insulated Winter Jacket', 'Warm synthetic insulation that still works when wet.', ['Rated to -10 C', 'Water-repellent', 'Sizes XS-XXL']),
        ('Lightweight Windbreaker', 'A light layer for windy runs and walks that packs into its own pocket.', ['Weighs 180 g', 'Reflective trims', 'Sizes S-XL']),
    ]),
    'Shoes': ('Apparel', ['Stride', 'Northtrail'], (49, 179), [
        ('Lightweight Running Shoes', 'Cushioned, breathable running shoes for daily training.', ['Weighs {grams} g', 'Breathable mesh', 'Sizes 5-12']),
        ('Waterproof Trail Shoes', 'Grippy outsole and waterproof membrane for muddy trails.', ['Waterproof membrane', 'Rock plate', 'Sizes 5-12']),
        ('Everyday Sneakers', 'Clean, comfortable sneakers for the office and weekends.', ['Leather upper', 'Removable insole', 'Sizes 5-12']),
    ]),
    'Kitchen': ('Home & Kitchen', ['Hearth', 'Brewly'], (15, 199), [
        ('Pour-Over Coffee Set', 'Glass dripper, carafe and filters for a clean, bright cup.', ['600 ml carafe', '40 filters included', 'Dishwasher safe']),
        ('Insulated Steel Bottle', 'Keeps drinks cold for 24 hours or hot for 12.', ['750 ml', 'Leak-proof lid', 'BPA free']),
        ('Cast Iron Skillet', 'Pre-seasoned skillet that goes from stove to oven.', ['26 cm', 'Oven safe', 'Lifetime durability']),
        ('Electric Kettle', 'Fast-boil kettle with temperature settings for tea and coffee.', ['1.7 litre', '5 temperature presets', 'Keep-warm mode']),
    ]),
    'Fitness': ('Fitness', ['FlexFit', 'Corestrong'], (15, 299), [
        ('Yoga Mat', 'Non-slip mat with enough cushioning for knees and wrists.', ['6 mm thick', 'Carry strap', 'Non-slip both sides']),
        ('Adjustable Dumbbells', 'Change weight in seconds with a turn of the dial.', ['2.5-24 kg each', 'Space saving', 'Storage tray']),
        ('Resistance Band Set', 'Five bands of different strength for home workouts.', ['5 levels', 'Door anchor', 'Exercise guide']),
        ('Fitness Tracker', 'Tracks steps, heart rate and sleep, with a {days}-day battery.', ['Heart-rate sensor', 'Sleep tracking', 'Water resistant 50 m']),
    ]),
}
COLOURS = ['Black', 'White', 'Grey', 'Navy', 'Green', 'Sand', 'Red']


def make_catalog():
    rows = []
    pid = 1000
    for group, (category, brands, (low, high), templates) in CATALOG_SPEC.items():
        for brand in brands:
            for name, desc, features in templates:
                for tier, mult in (('', 1.0), (' Pro', 1.45)):
                    if tier and rng.random() < 0.45:
                        continue
                    pid += 1
                    base = rng.uniform(low, high * 0.7)
                    price = round(min(high, base * mult)) - 0.01
                    text = desc.format(hours=rng.choice([20, 24, 30, 40]), width=rng.choice([90, 100, 110]),
                                       grams=rng.choice([220, 240, 260]), days=rng.choice([7, 10, 14]))
                    feats = [f.format(width=rng.choice([90, 100, 110]), grams=rng.choice([220, 240, 260]),
                                      days=rng.choice([7, 10, 14])) for f in features]
                    if tier:
                        text += ' The Pro version adds premium materials and a 2-year warranty.'
                    colours = ', '.join(rng.sample(COLOURS, rng.randint(1, 3)))
                    stock = rng.choices([0, rng.randint(1, 8), rng.randint(9, 120)], weights=[0.1, 0.2, 0.7])[0]
                    rating = round(rng.uniform(3.6, 4.9), 1)
                    rows.append([f'P{pid}', f'{brand} {name}{tier}', category, group, brand, f'{price:.2f}',
                                 colours, text, '; '.join(feats), stock, 'In stock' if stock > 8 else ('Low stock' if stock else 'Out of stock'),
                                 rating, rng.randint(4, 900)])
    write_csv(ROOT / 'shopassist' / 'products.csv',
              ['product_id', 'name', 'category', 'subcategory', 'brand', 'price_usd', 'colours', 'description',
               'features', 'stock_qty', 'availability', 'rating', 'review_count'], rows)


# ---------------------------------------------------------------------------
# 3. BookIt: sports venue bookings (past 120 days + next 14 days)
# ---------------------------------------------------------------------------
def make_bookings():
    rows = []
    courts = [('Turf A (5-a-side)', 60), ('Turf B (5-a-side)', 60), ('Turf C (7-a-side)', 90), ('Badminton Court 1', 15), ('Badminton Court 2', 15)]
    today = date(2026, 10, 22)
    bid = 5000
    for d in daterange(today - timedelta(days=120), today + timedelta(days=14)):
        for court, rate in courts:
            hours = range(6, 23)
            for h in hours:
                evening = h >= 17
                weekend = d.weekday() >= 5
                busy = 0.25 + (0.45 if evening else 0) + (0.2 if weekend else 0)
                if d > today:
                    busy *= 0.55 * (1 - (d - today).days / 20)  # future slots fill closer to the day
                if rng.random() > busy:
                    continue
                bid += 1
                price = rate * (1.3 if evening or weekend else 1.0)
                if d > today:
                    status = 'Confirmed'
                else:
                    status = rng.choices(['Completed', 'Cancelled', 'No-show'], weights=[0.86, 0.1, 0.04])[0]
                rows.append([f'BK-{bid}', d.isoformat(), f'{h:02d}:00', f'{h + 1:02d}:00', court,
                             f'Customer C{rng.randint(1000, 1999)}', rng.choice(['Website', 'Website', 'Phone', 'Walk-in']),
                             status, f'{price:.2f}', 'Yes' if rng.random() < 0.3 else 'No'])
    write_csv(ROOT / 'bookit' / 'bookings.csv',
              ['booking_id', 'date', 'start_time', 'end_time', 'resource', 'customer_ref', 'channel',
               'status', 'price_usd', 'equipment_rental'], rows)


# ---------------------------------------------------------------------------
# 4. SupportFlow: labelled customer messages with agent replies
# ---------------------------------------------------------------------------
TICKET_TEMPLATES = [
    ('Returns', 'Urgent', 'My {item} arrived with a cracked {part}. Can I get a replacement?',
     'We are sorry your {item} arrived damaged. Please reply with a photo of the cracked {part} and we will send a replacement within 1 working day. There is no need to return the damaged one.'),
    ('Returns', 'Urgent', 'I received a {wrong} instead of the {item} I ordered (order {order}).',
     'Sorry for the mix-up with order {order}. We will send the correct {item} today and email you a free return label for the {wrong}.'),
    ('Returns', 'Normal', 'I would like to return the {item}, it does not fit my needs. It is unused.',
     'No problem. Unused items can be returned within 14 days of delivery. Use the return link in your order email and your refund will reach you within 5 to 7 working days after we receive it.'),
    ('Returns', 'Normal', 'Can I exchange the {item} for a different colour?',
     'Yes, exchanges are free within 14 days. Start an exchange from your order page and choose the new colour; we ship it as soon as the original is on its way back.'),
    ('Delivery', 'Urgent', 'Order {order} was due {days} days ago and still has not arrived. I need it by {weekday}.',
     'Sorry for the delay with order {order}. We have opened an investigation with the courier and will update you within 24 hours. If it cannot reach you by {weekday}, we will ship a replacement by express at no cost.'),
    ('Delivery', 'Normal', 'Can I change the delivery address for order {order}?',
     'Yes, as long as the order has not been dispatched. Please reply with the new address and we will update order {order} right away.'),
    ('Delivery', 'Low', 'How long does standard delivery usually take?',
     'Orders are dispatched within 1 working day. Standard delivery takes 3 to 5 working days and express takes 1 to 2.'),
    ('Delivery', 'Low', 'Do you ship to other countries?',
     'At the moment we only ship within the country. We will announce on our website when international shipping starts.'),
    ('Payment', 'Urgent', 'I was charged twice for order {order}. Please fix this.',
     'Sorry about the double charge on order {order}. The duplicate payment is refunded automatically within 5 working days. We have flagged it so you do not need to do anything else.'),
    ('Payment', 'Normal', 'My card payment failed but the money left my account.',
     'Failed payments are reversed by your bank, usually within 3 to 5 working days. If it has not returned after that, reply with the date and amount and we will raise it with our payment provider.'),
    ('Payment', 'Low', 'Which payment methods do you accept?',
     'We accept cards and UPI. Cash on delivery is not available at the moment.'),
    ('Account', 'Normal', 'I cannot log in to my account, the password reset email never arrives.',
     'Please check your spam folder for an email from our store. If it is not there, reply with the email address on your account and we will resend it. We never ask for your password by email or chat.'),
    ('Account', 'Low', 'How do I change the email address on my account?',
     'Go to Account, then Profile, and choose Change email. You will get a confirmation link at the new address.'),
    ('Product question', 'Normal', 'Is the {item} compatible with {device}?',
     'Yes, the {item} works with {device}. If you have any trouble setting it up, our setup guide is on the product page.'),
    ('Product question', 'Low', 'What is the warranty on the {item}?',
     'The {item} has a 1-year warranty, or 2 years for Pro models. Keep your order email as proof of purchase.'),
    ('Product question', 'Normal', 'When will the {item} be back in stock?',
     'We expect new stock of the {item} within 2 weeks. Use the Notify me button on the product page and we will email you the moment it is back.'),
    ('Other', 'Low', 'Just wanted to say the {item} is great, thank you!',
     'Thank you for the kind words! We are glad you are enjoying the {item}.'),
    ('Other', 'Urgent', 'The {item} charger got very hot and smelled of burning.',
     'Please stop using the charger and unplug it now. Your safety comes first. We will send a replacement immediately and arrange collection of the faulty one; please reply with your order number.'),
]
ITEMS = [('headphones', 'ear cup'), ('webcam', 'lens'), ('desk lamp', 'base'), ('kettle', 'lid'), ('speaker', 'grille'),
         ('keyboard', 'key'), ('water bottle', 'cap'), ('fitness tracker', 'screen'), ('yoga mat', 'edge')]
DEVICES = ['a MacBook', 'Windows 11', 'an iPad', 'an Android phone', 'a Chromebook']
WEEKDAYS = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']
OPENERS = ['', 'Hi, ', 'Hello, ', 'Good morning. ', 'Hey team, ']
CLOSERS = ['', ' Thanks.', ' Thank you.']


def make_tickets():
    rows = []
    start = datetime(2026, 6, 1, 8, 0)
    for n in range(1, 801):
        topic, urgency, msg, reply = rng.choice(TICKET_TEMPLATES)
        item, part = rng.choice(ITEMS)
        wrong = rng.choice([i for i, _ in ITEMS if i != item])
        values = dict(item=item, part=part, wrong=wrong, order=f'#{rng.randint(40000, 49999)}',
                      days=rng.randint(3, 9), weekday=rng.choice(WEEKDAYS), device=rng.choice(DEVICES))
        body = msg.format(**values)
        opener = rng.choice(OPENERS)
        if opener in ('Hi, ', 'Hello, ', 'Hey team, ') and not body.startswith('I '):
            body = body[0].lower() + body[1:]
        closer = rng.choice(CLOSERS) if topic != 'Other' else ''
        text = opener + body + closer
        created = start + timedelta(minutes=rng.randint(0, 140 * 24 * 60))
        first_reply_mins = rng.randint(5, 90) if urgency == 'Urgent' else rng.randint(30, 600)
        rows.append([f'T-{n:05d}', created.strftime('%Y-%m-%d %H:%M'), rng.choice(['Email', 'Chat', 'Contact form']),
                     text, topic, urgency, reply.format(**values) + ' - The Support Team', first_reply_mins,
                     rng.choices(['Resolved', 'Resolved', 'Resolved', 'Waiting on customer'], k=1)[0]])
    write_csv(ROOT / 'supportflow' / 'tickets.csv',
              ['ticket_id', 'created_at', 'channel', 'message', 'topic', 'urgency', 'agent_reply',
               'first_reply_minutes', 'status'], rows)


# ---------------------------------------------------------------------------
# 5. ForecastBoard: weekly sales by category (about 2.75 years)
# ---------------------------------------------------------------------------
def make_weekly_sales():
    rows = []
    categories = {'Audio': (420, 89), 'Computer Accessories': (510, 64), 'Office Furniture': (95, 210),
                  'Apparel': (380, 58), 'Home & Kitchen': (300, 38), 'Fitness': (210, 55)}
    week = date(2024, 1, 1)
    holiday_weeks = {(11, 4), (12, 3), (12, 4), (1, 1), (8, 3)}  # (month, week-of-month)
    while week <= date(2026, 9, 28):
        t = (week - date(2024, 1, 1)).days / 7
        wom = (week.day - 1) // 7 + 1
        holiday = (week.month, wom) in holiday_weeks
        for cat, (base, price) in categories.items():
            season = seasonal(week + timedelta(days=3)) / 1.05
            if cat == 'Apparel':
                season *= 1 + 0.35 * math.cos((week.month - 12) / 12 * 2 * math.pi)
            if cat == 'Fitness' and week.month == 1:
                season *= 1.7
            promo = 1 if rng.random() < 0.08 else 0
            units = base * season * (1 + 0.0025 * t) * (1.25 if promo else 1) * (1.2 if holiday else 1)
            units = max(0, int(rng.gauss(units, units * 0.08)))
            avg_price = round(price * (0.88 if promo else 1) * rng.uniform(0.97, 1.03), 2)
            rows.append([week.isoformat(), cat, units, f'{avg_price:.2f}', promo, 1 if holiday else 0,
                         f'{units * avg_price:.2f}'])
        week += timedelta(days=7)
    write_csv(ROOT / 'forecastboard' / 'weekly_sales.csv',
              ['week_start', 'category', 'units_sold', 'avg_price', 'promotion', 'holiday_week', 'revenue'], rows)


# ---------------------------------------------------------------------------
# 6. VoiceDesk: clinic services and fees
# ---------------------------------------------------------------------------
def make_clinic_services():
    services = [
        ('General consultation', 'Consultations', 15, 500, 'None', 'Walk-in or appointment'),
        ('Follow-up consultation (within 7 days)', 'Consultations', 10, 0, 'Bring previous reports', 'Appointment'),
        ('Child health check', 'Child health', 20, 600, 'Bring vaccination card', 'Appointment'),
        ('Childhood vaccination', 'Vaccinations', 10, 300, 'Bring vaccination card', 'Appointment'),
        ('Adult flu vaccine', 'Vaccinations', 10, 900, 'None', 'Walk-in or appointment'),
        ('Travel vaccination advice', 'Vaccinations', 20, 700, 'Bring travel dates', 'Appointment'),
        ('Complete blood count', 'Lab tests', 5, 350, 'No fasting needed', 'Walk-in 8:00-11:00'),
        ('Blood sugar (fasting)', 'Lab tests', 5, 150, 'Fast for 10 hours', 'Walk-in 8:00-11:00'),
        ('Cholesterol (lipid profile)', 'Lab tests', 5, 650, 'Fast for 10 hours', 'Walk-in 8:00-11:00'),
        ('Thyroid profile', 'Lab tests', 5, 550, 'No fasting needed', 'Walk-in 8:00-11:00'),
        ('Vitamin D test', 'Lab tests', 5, 1200, 'No fasting needed', 'Walk-in 8:00-11:00'),
        ('Annual health check package', 'Packages', 60, 2900, 'Fast for 10 hours', 'Appointment'),
        ('Women\'s health check', 'Packages', 45, 2400, 'Fast for 10 hours', 'Appointment'),
        ('Wound dressing', 'Minor procedures', 15, 250, 'None', 'Walk-in'),
        ('Stitch removal', 'Minor procedures', 10, 300, 'None', 'Walk-in'),
        ('Ear cleaning', 'Minor procedures', 15, 400, 'Use ear drops for 3 days before', 'Appointment'),
        ('ECG', 'Diagnostics', 15, 450, 'None', 'Appointment'),
        ('Blood pressure check', 'Diagnostics', 5, 0, 'None', 'Walk-in'),
        ('Diet consultation', 'Wellness', 30, 800, 'Bring a 3-day food diary', 'Appointment'),
        ('Physiotherapy session', 'Wellness', 45, 900, 'Wear comfortable clothes', 'Appointment'),
        ('Home sample collection', 'Lab tests', 20, 200, 'Book by 18:00 the day before', 'Appointment'),
        ('Medical certificate', 'Documents', 10, 300, 'Needs a consultation first', 'Appointment'),
    ]
    write_csv(ROOT / 'voicedesk' / 'services.csv',
              ['service', 'category', 'duration_minutes', 'fee', 'preparation', 'how_to_book'], services)


# ---------------------------------------------------------------------------
# 7. ReportBot: cafe workbook with Transactions, Targets and Expenses sheets
# ---------------------------------------------------------------------------
def make_cafe_workbook():
    menu = [('Espresso', 'Coffee', 2.5), ('Cappuccino', 'Coffee', 3.6), ('Latte', 'Coffee', 3.9), ('Cold Brew', 'Coffee', 4.2),
            ('Masala Chai', 'Tea', 2.8), ('Green Tea', 'Tea', 2.4), ('Croissant', 'Bakery', 2.9), ('Blueberry Muffin', 'Bakery', 3.1),
            ('Banana Bread', 'Bakery', 3.3), ('Chicken Sandwich', 'Food', 6.9), ('Veg Wrap', 'Food', 6.2), ('Caesar Salad', 'Food', 7.5),
            ('Fresh Orange Juice', 'Cold drinks', 4.5), ('Iced Lemon Tea', 'Cold drinks', 3.4)]
    outlets = {'City Centre': 1.3, 'Tech Park': 1.1, 'Lakeside': 0.8}
    wb = Workbook()
    tx = wb.active
    tx.title = 'Transactions'
    tx.append(['date', 'outlet', 'item', 'category', 'quantity', 'unit_price', 'revenue', 'payment_method', 'order_type'])
    monthly = {}
    for d in daterange(date(2025, 10, 1), date(2026, 9, 30)):
        for outlet, w in outlets.items():
            wk = 1.15 if (outlet == 'Lakeside' and d.weekday() >= 5) else (0.6 if outlet == 'Tech Park' and d.weekday() >= 5 else 1)
            for item, cat, price in menu:
                hot = cat in ('Coffee', 'Tea')
                temp_factor = 1.25 if (hot and d.month in (11, 12, 1, 2)) else (1.4 if cat == 'Cold drinks' and d.month in (4, 5, 6) else 1)
                qty = max(0, int(rng.gauss(9, 3) * w * wk * temp_factor))
                if not qty:
                    continue
                rev = round(qty * price, 2)
                tx.append([d, outlet, item, cat, qty, price, rev, rng.choice(['Card', 'UPI', 'Cash']),
                           rng.choices(['Dine-in', 'Takeaway', 'Delivery'], weights=[0.5, 0.35, 0.15])[0]])
                key = (d.strftime('%Y-%m'), outlet)
                monthly[key] = monthly.get(key, 0) + rev
    tg = wb.create_sheet('Targets')
    tg.append(['month', 'outlet', 'revenue_target'])
    for (month, outlet), actual in sorted(monthly.items()):
        tg.append([month, outlet, round(actual * rng.uniform(0.92, 1.1) / 100) * 100])
    ex = wb.create_sheet('Expenses')
    ex.append(['month', 'outlet', 'expense_type', 'amount'])
    for (month, outlet), actual in sorted(monthly.items()):
        for etype, share in (('Ingredients', 0.31), ('Staff', 0.27), ('Rent', 0.12), ('Utilities', 0.04), ('Marketing', 0.02)):
            ex.append([month, outlet, etype, round(actual * share * rng.uniform(0.9, 1.1), 2)])
    for ws in (tx, tg, ex):
        for cell in ws[1]:
            cell.font = Font(bold=True)
    for row in tx.iter_rows(min_row=2, max_col=1):
        row[0].number_format = 'yyyy-mm-dd'
    path = ROOT / 'reportbot' / 'cafe_sales.xlsx'
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f'{path.relative_to(ROOT)}: {tx.max_row - 1} transactions, {tg.max_row - 1} targets, {ex.max_row - 1} expenses')


if __name__ == '__main__':
    make_retail_sales()
    make_catalog()
    make_bookings()
    make_tickets()
    make_weekly_sales()
    make_clinic_services()
    make_cafe_workbook()
